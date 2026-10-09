import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, Alignment
from natsort import natsorted, natsort_keygen
from pyomo.environ import *

# ---------------------------
# READ EXCEL DATA
# ---------------------------
file = "school_data_realistic.xlsx"
qual_df = pd.read_excel(file, sheet_name="qualifications").fillna(0)
pref_df = pd.read_excel(file, sheet_name="preferences").fillna(0)
work_df = pd.read_excel(file, sheet_name="workload").fillna(0)

# ---------------------------
# CONVERT DATAFRAMES TO DICTIONARIES
# ---------------------------
qual_data = {(r.Teacher, r.Subject): int(r.Qualified) for _, r in qual_df.iterrows()}
pref_data = {(r.Teacher, r.Period): int(r.Preference) for _, r in pref_df.iterrows()}
work_data = {(r.Class, r.Subject): int(r.Periods) for _, r in work_df.iterrows()}

# ---------------------------
# CREATE PYOMO SETS
# ---------------------------
teachers = natsorted(list(set(qual_df["Teacher"])))
subjects = natsorted(list(set(qual_df["Subject"])))
classes = natsorted(list(set(work_df["Class"])))
periods = natsorted(list(set(pref_df["Period"])))

# ---------------------------
# CREATE CONCRETE MODEL
# ---------------------------
model = ConcreteModel()

# ---------------------------
# DEFINE SETS
# ---------------------------
model.T = Set(initialize=teachers)
model.S = Set(initialize=subjects)
model.C = Set(initialize=classes)
model.P = Set(initialize=periods)

# ---------------------------
# DEFINE PARAMETERS
# ---------------------------
model.Q = Param(model.T, model.S, initialize=qual_data, within=Binary, default=0)
model.Pref = Param(model.T, model.P, initialize=pref_data, within=NonNegativeIntegers, default=0)
model.W = Param(model.C, model.S, initialize=work_data, within=NonNegativeIntegers, default=0)

# ---------------------------
# DECISION VARIABLES
# ---------------------------

# x[t,c,s,p] = 1 if teacher t teaches class c, subject s, period p
model.x = Var(model.T, model.C, model.S, model.P, domain=Binary)

# y[t,c,s] = 1 if teacher t is assigned to teach all periods of class c, subject s
model.y = Var(model.T, model.C, model.S, domain=Binary)

# workload per teacher
model.workload = Var(model.T, domain=NonNegativeIntegers)

# workload deviation per teacher
model.dev = Var(model.T, domain=NonNegativeReals)

# ---------------------------
# CONSTRAINTS
# ---------------------------

# 1. Teacher qualification
def qualification_rule(model, t, c, s, p):
    return model.x[t, c, s, p] <= model.Q[t, s]
model.Qualified = Constraint(model.T, model.C, model.S, model.P, rule=qualification_rule)

# 2. One teacher per class-subject
def single_teacher_class_subject_rule(model, c, s):
    return sum(model.y[t, c, s] for t in model.T) == 1
model.SingleTeacher = Constraint(model.C, model.S, rule=single_teacher_class_subject_rule)

# 3. Link x and y to enforce consistency
def consistent_teacher_rule(model, t, c, s, p):
    return model.x[t, c, s, p] <= model.y[t, c, s]
model.ConsistentTeacher = Constraint(model.T, model.C, model.S, model.P, rule=consistent_teacher_rule)

# 4. No double booking
def no_double_booking_rule(model, t, p):
    return sum(model.x[t, c, s, p] for c in model.C for s in model.S) <= 1
model.NoDoubleBooking = Constraint(model.T, model.P, rule=no_double_booking_rule)

# 5. Workload requirement
def workload_requirement_rule(model, c, s):
    return sum(model.x[t, c, s, p] for t in model.T for p in model.P) == model.W[c, s]
model.WorkloadRequirement = Constraint(model.C, model.S, rule=workload_requirement_rule)

# 6. Compute teacher workload
def workload_rule(model, t):
    return model.workload[t] == sum(model.x[t, c, s, p] for c in model.C for s in model.S for p in model.P)
model.TeacherWorkload = Constraint(model.T, rule=workload_rule)

# ---------------------------------
# BI-CRITERIA OBJECTIVE FUNCTION
# ----------------------------------

# Maximize preference satisfaction
pref_expr = sum(model.Pref[t, p]*model.x[t, c, s, p] for t in model.T for c in model.C for s in model.S for p in model.P)

# Minimize workload imbalance (linearized using deviations from mean)
mean_workload = sum(model.W[c,s] for c in model.C for s in model.S)/len(model.T)  # approximate mean

def abs_dev_pos_rule(model, t):
    return model.dev[t] >= model.workload[t] - mean_workload
model.AbsDevPos = Constraint(model.T, rule=abs_dev_pos_rule)

def abs_dev_neg_rule(model, t):
    return model.dev[t] >= mean_workload - model.workload[t]
model.AbsDevNeg = Constraint(model.T, rule=abs_dev_neg_rule)

workload_imbalance_expr = sum(model.dev[t] for t in model.T)

alpha = 1.0
beta = 0.1

model.Objective = Objective(expr = alpha*pref_expr - beta*workload_imbalance_expr, sense=maximize)

#----------------------------------------------
# TIME SLOT MAPPING  
#-------------------------------

def time_slot_mapping():
    timeslot_map = {}
    periods_per_day = 7
    days = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"]

    for i in range(1, 36):  # P1–P35
        P = f"P{i}"
        day_index = (i - 1) // periods_per_day
        period_in_day = (i - 1) % periods_per_day + 1

        day = days[day_index]
        timeslot_map[P] = f"{day}, Period {period_in_day}"

    return timeslot_map


timeslot_map = time_slot_mapping()

#---------------------------------------
# BUILDING THE ASSIGNMENT DATAFRAME
#------------------------------------------

def assignment_df(model, teachers, classes, subjects, periods, timeslot_map):
    rows = []

    for t in teachers:
        for c in classes:
            for s in subjects:
                for p in periods:
                    if model.x[t,c,s,p].value == 1:
                        period_str = str(p).strip()
                        ts = timeslot_map[period_str]
                        rows.append({
                            "Teacher": t, 
                            "Class": c, 
                            "Subject": s, 
                            "Period": period_str,
                            "Timeslot": ts

                        })

    df = pd.DataFrame(rows)

    df = df.sort_values(["Teacher", "Period"], key=lambda col: (col.str.extract(r"(\d+)").astype(int)[0]
        if col.name in ["Teacher", "Period"]      
        else col

    )       

)           
    return df



# ---------------------------
# SOLVE MILP MODEL
# ---------------------------
solver = SolverFactory('highs')
#solver.options['mip_rel_gap'] = 5
results = solver.solve(model, tee=True)

# ---------------------------
# DISPLAY RESULTS
# ---------------------------
eps = 1e-6
for t in model.T:
    for c in model.C:
        for s in model.S:
            for p in model.P:
                if value(model.x[t,c,s,p]) > 1-eps:
                    print(f"Teacher {t} teaches {s} to {c} in Period {p}")

print("\nWorkloads")
for t in model.T:
    print(f"{t}: {value(model.workload[t])}")


assign_df = assignment_df(
    model,
    teachers=list(model.T),
    classes=list(model.C),
    subjects=list(model.S),
    periods = list(model.P),
    timeslot_map=timeslot_map
                          
)

print(assign_df)

# ==========================================================
# 1. MASTER TIMETABLE (already built from assignment_df())
# ==========================================================

master_df = assign_df.copy()

# ==========================================================
# 2. TEACHER WORKLOAD SUMMARY
# ==========================================================

workload_rows = []
for t in model.T:
    workload_rows.append({
        "Teacher": t,
        "Workload": value(model.workload[t]),
        "Deviation": value(model.dev[t])
    })


workload_df = (
    pd.DataFrame(workload_rows)
    .sort_values(
        "Teacher",
        key=lambda col: col.str.extract(r"(\d+)").astype(int)[0]
    )
)



# ==========================================================
# 3. CLASS–SUBJECT RESPONSIBLE TEACHER
# ==========================================================

cs_rows = []
for c in model.C:
    for s in model.S:
        for t in model.T:
            if value(model.y[t, c, s]) > 1-eps:
                cs_rows.append({
                    "Class": c,
                    "Subject": s,
                    "Teacher": t
                })

class_subject_df = pd.DataFrame(cs_rows).sort_values(["Class", "Subject"])


# ==========================================================
# 4. TEACHER TIMETABLES (Pivot format)
# ==========================================================


master_df["Day"] = master_df["Timeslot"].str.split(",").str[0]

master_df["PeriodInDay"] = (
    master_df["Timeslot"]
    .str.extract(r"Period (\d+)")[0]
    .astype(int)
)



teacher_tables = (
    master_df
    .pivot_table(
        index=["Teacher", "Day"],
        columns="PeriodInDay",
        values="Subject",
        aggfunc=lambda x: ", ".join(x)
    )
    .fillna("")
    .reset_index()
)

day_order = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"]
teacher_order = [f"T{i}" for i in range(1, 11)]

teacher_tables["Day"] = pd.Categorical(teacher_tables["Day"], categories=day_order, ordered=True)

teacher_tables["Teacher"] = pd.Categorical(teacher_tables["Teacher"], categories=teacher_order, ordered=True)

teacher_tables = teacher_tables.sort_values(["Teacher", "Day"])



# ==========================================================
# 5. RAW X VARIABLE ASSIGNMENTS
# ==========================================================

raw_rows = []
for t in model.T:
    for c in model.C:
        for s in model.S:
            for p in model.P:
                if value(model.x[t, c, s, p]) > 0.5:
                    raw_rows.append({
                        "Teacher": t,
                        "Class": c,
                        "Subject": s,
                        "Period": p
                    })

raw_x_df = pd.DataFrame(raw_rows)

# ==========================================================
# EXPORT TO MULTI-SHEET EXCEL FILE
# ==========================================================

output_file = "Full_Timetable_Report.xlsx"

with pd.ExcelWriter(output_file, engine="openpyxl") as writer:
    master_df.to_excel(writer, sheet_name="Master Timetable", index=False)
    workload_df.to_excel(writer, sheet_name="Workloads", index=False)
    class_subject_df.to_excel(writer, sheet_name="Class-Subject", index=False)
    teacher_tables.to_excel(writer, sheet_name="Teacher Timetables", index=False)
    raw_x_df.to_excel(writer, sheet_name="Raw Assignments", index=False)

# ==========================================================
# OPTIONAL: AUTO-FORMAT EXCEL FILE
# ==========================================================

wb = load_workbook(output_file)
for sheet in wb.sheetnames:
    ws = wb[sheet]

    # Bold header row + center alignment
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    # Autofit columns
    for col_cells in ws.columns:
        length = max(len(str(cell.value)) for cell in col_cells)
        ws.column_dimensions[get_column_letter(col_cells[0].column)].width = length + 2

wb.save(output_file)

print(f"\nFull reporting file exported to: {output_file}")
